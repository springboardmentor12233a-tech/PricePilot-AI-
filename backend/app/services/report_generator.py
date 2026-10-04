"""
Business Intelligence Report Generator (PDF, CSV, and Excel).
Generates an executive-ready PDF report using ReportLab with embedded Matplotlib charts,
a structured CSV export using pandas, and a multi-sheet formatted Excel workbook.
"""

import io
from datetime import datetime
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.services.data_loader import get_all_products, get_product_by_id, calculate_kpi_overview, get_revenue_history
from app.ml.model_registry import registry


def _render_revenue_chart_image() -> io.BytesIO:
    history = get_revenue_history(days=90)
    months = [h["month"].split()[0] for h in history[-8:]]
    baseline = [h["baseline_revenue"] / 1000 for h in history[-8:]]
    optimized = [h["optimized_revenue"] / 1000 for h in history[-8:]]

    fig, ax = plt.subplots(figsize=(6.5, 2.4), dpi=180)
    fig.patch.set_facecolor("#0F172A")
    ax.set_facecolor("#1E293B")

    ax.plot(months, baseline, color="#94A3B8", linestyle="--", marker="o", linewidth=1.8, label="Baseline ($K)")
    ax.plot(months, optimized, color="#14B8A6", linestyle="-", marker="s", linewidth=2.4, label="PricePilot Optimized ($K)")
    ax.fill_between(months, baseline, optimized, color="#14B8A6", alpha=0.15)

    ax.set_title("Portfolio Revenue Performance: Baseline vs. Optimized (in Thousands USD)", color="#F8FAFC", fontsize=9, fontweight="bold", pad=8)
    ax.tick_params(colors="#94A3B8", labelsize=8)
    for spine in ax.spines.values():
        spine.set_color("#334155")
    ax.grid(color="#334155", linestyle=":", alpha=0.6)
    ax.legend(facecolor="#1E293B", edgecolor="#334155", labelcolor="#F8FAFC", fontsize=8)

    buf = io.BytesIO()
    plt.tight_layout()
    plt.savefig(buf, format="png", facecolor=fig.get_facecolor(), edgecolor="none")
    plt.close(fig)
    buf.seek(0)
    return buf


def generate_pdf_report(user_email: str = "analyst@pricepilot.ai") -> io.BytesIO:
    products = get_all_products()
    kpis = calculate_kpi_overview()
    meta = registry.get_metadata()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    primary_color = colors.HexColor("#0F172A")
    teal_accent = colors.HexColor("#0D9488")
    slate_muted = colors.HexColor("#475569")
    card_bg = colors.HexColor("#F8FAFC")
    border_color = colors.HexColor("#E2E8F0")

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=primary_color,
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=slate_muted,
    )
    section_heading = ParagraphStyle(
        "SectionHeading",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=teal_accent,
        spaceBefore=8,
        spaceAfter=5,
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=10,
        textColor=primary_color,
    )
    bold_body = ParagraphStyle("BoldBody", parent=body_style, fontName="Helvetica-Bold")

    elements = []

    # 1. Header Banner
    header_data = [
        [
            Paragraph("<b>PRICEPILOT AI</b> &bull; REVENUE INTELLIGENCE REPORT", ParagraphStyle("H1", fontName="Helvetica-Bold", fontSize=13, textColor=teal_accent)),
            Paragraph(f"Generated: {datetime.utcnow().strftime('%b %d, %Y %H:%M UTC')}<br/>Authorized User: {user_email}", ParagraphStyle("H2", fontName="Helvetica", fontSize=7.5, alignment=2, textColor=slate_muted)),
        ]
    ]
    header_table = Table(header_data, colWidths=[340, 200])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(header_table)
    elements.append(HRFlowable(width="100%", thickness=1.5, color=teal_accent, spaceBefore=4, spaceAfter=8))

    elements.append(Paragraph("Executive Pricing & Revenue Intelligence Summary", title_style))
    elements.append(Paragraph(
        f"Econometric Model: beta = {meta['elasticity_coefficient']:.3f}, R² = {meta['r_squared']:.3f}, p = {meta['p_value']:.4f} | Electronics Portfolio",
        subtitle_style
    ))
    elements.append(Spacer(1, 8))

    # 2. KPI Summary Cards Grid
    kpi_table_data = [
        [
            Paragraph("<b>Total Monthly Revenue</b>", subtitle_style),
            Paragraph("<b>Revenue Growth MoM</b>", subtitle_style),
            Paragraph("<b>Potential Revenue Lift</b>", subtitle_style),
            Paragraph("<b>Model Confidence (Avg)</b>", subtitle_style),
        ],
        [
            Paragraph(f"<b>${kpis['total_monthly_revenue']:,.2f}</b>", ParagraphStyle("KPI1", fontName="Helvetica-Bold", fontSize=12, textColor=primary_color)),
            Paragraph(f"<b>{kpis['revenue_growth_pct']:+.1f}%</b>", ParagraphStyle("KPI2", fontName="Helvetica-Bold", fontSize=12, textColor=teal_accent)),
            Paragraph(f"<b>+${kpis['potential_revenue_lift']:,.2f}</b>", ParagraphStyle("KPI3", fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor("#D97706"))),
            Paragraph(f"<b>{kpis['model_avg_confidence']:.1f}%</b>", ParagraphStyle("KPI4", fontName="Helvetica-Bold", fontSize=12, textColor=primary_color)),
        ],
    ]
    kpi_table = Table(kpi_table_data, colWidths=[135, 135, 135, 135])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), card_bg),
        ('BOX', (0,0), (-1,-1), 1, border_color),
        ('INNERGRID', (0,0), (-1,-1), 0.5, border_color),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 7),
        ('RIGHTPADDING', (0,0), (-1,-1), 7),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 8))

    # 3. Chart
    elements.append(Paragraph("Portfolio Revenue Trajectory & Optimization Lift", section_heading))
    chart_stream = _render_revenue_chart_image()
    chart_img = Image(chart_stream, width=7.4 * inch, height=2.4 * inch)
    elements.append(chart_img)
    elements.append(Spacer(1, 8))

    # 4. Top 3 AI Recommendations
    elements.append(Paragraph("Priority AI Pricing Recommendations", section_heading))
    top_products = sorted(products, key=lambda x: abs(x["price_gap_pct"]), reverse=True)[:3]
    rec_rows = [
        [
            Paragraph("<b>Product & Category</b>", bold_body),
            Paragraph("<b>Current</b>", bold_body),
            Paragraph("<b>Target</b>", bold_body),
            Paragraph("<b>Gap %</b>", bold_body),
            Paragraph("<b>AI Strategic Rationale</b>", bold_body),
        ]
    ]
    for tp in top_products:
        gap_color = "#DC2626" if tp["price_gap_pct"] < 0 else "#16A34A"
        rec_rows.append([
            Paragraph(f"<b>{tp['name']}</b><br/><font color='#64748B'>{tp['category']}</font>", body_style),
            Paragraph(f"${tp['current_price']:,.2f}", body_style),
            Paragraph(f"<b>${tp['recommended_price']:,.2f}</b>", body_style),
            Paragraph(f"<font color='{gap_color}'><b>{tp['price_gap_pct']:+.1f}%</b></font>", body_style),
            Paragraph(tp["llm_summary"], body_style),
        ])
    rec_table = Table(rec_rows, colWidths=[120, 50, 50, 50, 270])
    rec_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#E2E8F0")),
        ('BOX', (0,0), (-1,-1), 0.5, border_color),
        ('INNERGRID', (0,0), (-1,-1), 0.5, border_color),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('LEFTPADDING', (0,0), (-1,-1), 5),
        ('RIGHTPADDING', (0,0), (-1,-1), 5),
    ]))
    elements.append(rec_table)
    elements.append(Spacer(1, 8))

    # 5. Product Table
    elements.append(Paragraph("Complete Portfolio Pricing & Forecast Table", section_heading))
    prod_rows = [
        [
            Paragraph("<b>SKU Name</b>", bold_body),
            Paragraph("<b>Category</b>", bold_body),
            Paragraph("<b>Current</b>", bold_body),
            Paragraph("<b>Target</b>", bold_body),
            Paragraph("<b>Comp.</b>", bold_body),
            Paragraph("<b>Units</b>", bold_body),
            Paragraph("<b>Revenue</b>", bold_body),
            Paragraph("<b>Demand</b>", bold_body),
            Paragraph("<b>Conf.</b>", bold_body),
        ]
    ]
    for p in products:
        prod_rows.append([
            Paragraph(p["name"][:22] + "...", body_style),
            Paragraph(p["category"], body_style),
            Paragraph(f"${p['current_price']:.0f}", body_style),
            Paragraph(f"${p['recommended_price']:.0f}", body_style),
            Paragraph(f"${p['competitor_price']:.0f}", body_style),
            Paragraph(f"{p['units_sold']:,}", body_style),
            Paragraph(f"${p['revenue_this_month']/1000:.1f}k", body_style),
            Paragraph(p["demand_trend"][:8], body_style),
            Paragraph(f"{p['confidence_score']:.0f}%", body_style),
        ])
    prod_table = Table(prod_rows, colWidths=[120, 60, 48, 52, 55, 52, 55, 50, 48])
    prod_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0F172A")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('BOX', (0,0), (-1,-1), 0.5, border_color),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#F1F5F9")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, card_bg]),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(prod_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generate_csv_report() -> io.BytesIO:
    """
    Generates a structured CSV report using pandas.
    """
    products = get_all_products()
    df = pd.DataFrame(products)
    columns_to_export = [
        "id", "name", "category", "current_price", "recommended_price",
        "price_gap_pct", "cost_price", "competitor_price", "units_sold",
        "revenue_this_month", "predicted_revenue", "stock_status",
        "demand_trend", "confidence_score", "elasticity_coef"
    ]
    export_df = df[[c for c in columns_to_export if c in df.columns]]
    buffer = io.BytesIO()
    export_df.to_csv(buffer, index=False, encoding="utf-8")
    buffer.seek(0)
    return buffer


def generate_excel_report() -> io.BytesIO:
    wb = Workbook()
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    # Sheet 1: Products & Pricing
    ws1 = wb.active
    ws1.title = "Products & Pricing"
    ws1.views.sheetView[0].showGridLines = True

    headers = [
        "Product ID", "Product Name", "Category", "Current Price ($)",
        "Recommended Price ($)", "Price Gap (%)", "Competitor Price ($)",
        "Cost Price ($)", "Units Sold", "Revenue This Month ($)",
        "Predicted Revenue ($)", "Stock Status", "Demand Trend",
        "Confidence Score (%)", "Elasticity Beta"
    ]
    ws1.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for p in get_all_products():
        ws1.append([
            p["id"], p["name"], p["category"], p["current_price"],
            p["recommended_price"], p["price_gap_pct"], p["competitor_price"],
            p["cost_price"], p["units_sold"], p["revenue_this_month"],
            p["predicted_revenue"], p["stock_status"], p["demand_trend"],
            p["confidence_score"], p.get("elasticity_coef", registry.elasticity_coef)
        ])

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Sheet 2: Revenue Trajectory
    ws2 = wb.create_sheet(title="Revenue Performance")
    ws2.views.sheetView[0].showGridLines = True
    rev_headers = ["Month", "Baseline Revenue ($)", "Optimized Revenue ($)", "Lift ($)", "Units Sold"]
    ws2.append(rev_headers)
    for col_num in range(1, len(rev_headers) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in get_revenue_history(days=90):
        lift = r["optimized_revenue"] - r["baseline_revenue"]
        ws2.append([r["month"], r["baseline_revenue"], r["optimized_revenue"], lift, r["units_sold"]])

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 15)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _render_price_comparison_chart_image(product: dict) -> io.BytesIO:
    labels = ["Current", "Recommended", "Comp 1"]
    comp1 = product.get("comp_1") or product.get("competitor_price", 0.0)
    values = [product["current_price"], product["recommended_price"], comp1]
    colors_list = ["#94A3B8", "#14B8A6", "#F59E0B"]

    if product.get("comp_2"):
        labels.append("Comp 2")
        values.append(product["comp_2"])
        colors_list.append("#EAB308")
    if product.get("comp_3"):
        labels.append("Comp 3")
        values.append(product["comp_3"])
        colors_list.append("#FB923C")

    fig, ax = plt.subplots(figsize=(6.5, 2.6), dpi=180)
    fig.patch.set_facecolor("#0F172A")
    ax.set_facecolor("#1E293B")

    bars = ax.bar(labels, values, color=colors_list, width=0.45, edgecolor="#334155", linewidth=1)
    for bar in bars:
        height = bar.get_height()
        ax.annotate(
            f"${height:,.2f}",
            xy=(bar.get_x() + bar.get_width() / 2, height),
            xytext=(0, 4),
            textcoords="offset points",
            ha="center",
            va="bottom",
            color="#F8FAFC",
            fontsize=8,
            fontweight="bold",
        )

    ax.set_title(f"Price Benchmark Comparison - {product['name']}", color="#F8FAFC", fontsize=9, fontweight="bold", pad=10)
    ax.set_ylabel("Price ($ USD)", color="#94A3B8", fontsize=8)
    ax.tick_params(colors="#94A3B8", labelsize=8)
    for spine in ax.spines.values():
        spine.set_color("#334155")
    ax.grid(axis="y", color="#334155", linestyle=":", alpha=0.6)
    ax.set_ylim(0, max(values) * 1.15)

    buf = io.BytesIO()
    plt.tight_layout()
    plt.savefig(buf, format="png", facecolor=fig.get_facecolor(), edgecolor="none")
    plt.close(fig)
    buf.seek(0)
    return buf


def generate_price_comparison_pdf(product_id: str, user_email: str = "analyst@pricepilot.ai") -> io.BytesIO:
    product = get_product_by_id(product_id)
    if not product:
        prods = get_all_products()
        product = prods[0] if prods else {
            "id": product_id,
            "name": f"Product {product_id}",
            "category": "Electronics",
            "current_price": 199.99,
            "recommended_price": 189.99,
            "competitor_price": 209.99,
            "comp_1": 209.99,
            "comp_2": None,
            "comp_3": None,
            "cost_price": 120.0,
            "demand_trend": "Stable",
            "stock_status": "In Stock",
            "confidence_score": 32.8,
        }

    meta = registry.get_metadata()
    comp_price = product.get("comp_1") or product.get("competitor_price", product["current_price"])
    rec_price = product["recommended_price"]
    diff_amount = rec_price - comp_price
    diff_pct = (diff_amount / comp_price * 100) if comp_price > 0 else 0.0

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    primary_color = colors.HexColor("#0F172A")
    teal_accent = colors.HexColor("#0D9488")
    slate_muted = colors.HexColor("#475569")
    card_bg = colors.HexColor("#F8FAFC")
    border_color = colors.HexColor("#E2E8F0")

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=primary_color,
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=slate_muted,
    )
    section_heading = ParagraphStyle(
        "SectionHeading",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=teal_accent,
        spaceBefore=6,
        spaceAfter=4,
    )

    elements = []

    # 1. Header Banner
    header_data = [
        [
            Paragraph("<b>PRICEPILOT AI</b> &bull; SINGLE SKU PRICE COMPARISON", ParagraphStyle("H1", fontName="Helvetica-Bold", fontSize=12, textColor=teal_accent)),
            Paragraph(f"Exported: {datetime.utcnow().strftime('%b %d, %Y %H:%M UTC')}<br/>Requested by: {user_email}", ParagraphStyle("H2", fontName="Helvetica", fontSize=7.5, alignment=2, textColor=slate_muted)),
        ]
    ]
    header_table = Table(header_data, colWidths=[340, 200])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(header_table)
    elements.append(HRFlowable(width="100%", thickness=1.5, color=teal_accent, spaceBefore=4, spaceAfter=8))

    # 2. Product Title & Metadata
    elements.append(Paragraph(f"{product['name']}", title_style))
    elements.append(Paragraph(
        f"SKU ID: {product['id']} &bull; Category: {product['category']} &bull; Demand Trend: {product.get('demand_trend', 'Stable')} &bull; Model Confidence: {product.get('confidence_score', 32.8):.1f}%",
        subtitle_style
    ))
    elements.append(Spacer(1, 8))

    # 3. Price Comparison Metrics Table
    diff_color = "#10B981" if diff_amount < 0 else "#EF4444" if diff_amount > 0 else "#64748B"
    diff_sign = "+" if diff_amount > 0 else ""

    summary_data = [
        [
            Paragraph("<b>Current Price</b>", ParagraphStyle("H", fontName="Helvetica-Bold", fontSize=8, textColor=slate_muted)),
            Paragraph("<b>AI Recommended</b>", ParagraphStyle("H", fontName="Helvetica-Bold", fontSize=8, textColor=teal_accent)),
            Paragraph("<b>Benchmark (Comp 1)</b>", ParagraphStyle("H", fontName="Helvetica-Bold", fontSize=8, textColor=slate_muted)),
            Paragraph("<b>Rec vs. Comp Delta</b>", ParagraphStyle("H", fontName="Helvetica-Bold", fontSize=8, textColor=slate_muted)),
        ],
        [
            Paragraph(f"<b>${product['current_price']:,.2f}</b>", ParagraphStyle("V1", fontName="Helvetica-Bold", fontSize=12, textColor=primary_color)),
            Paragraph(f"<b>${rec_price:,.2f}</b>", ParagraphStyle("V2", fontName="Helvetica-Bold", fontSize=12, textColor=teal_accent)),
            Paragraph(f"<b>${comp_price:,.2f}</b>", ParagraphStyle("V3", fontName="Helvetica-Bold", fontSize=12, textColor=slate_muted)),
            Paragraph(f"<b>{diff_sign}${diff_amount:,.2f} ({diff_sign}{diff_pct:.1f}%)</b>", ParagraphStyle("V4", fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor(diff_color))),
        ]
    ]

    summary_table = Table(summary_data, colWidths=[135, 135, 135, 135])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), card_bg),
        ('BOX', (0,0), (-1,-1), 1, border_color),
        ('INNERGRID', (0,0), (-1,-1), 0.5, border_color),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 10))

    # 4. Competitor Benchmarks Detail
    elements.append(Paragraph("Competitor Price Benchmarks", section_heading))
    comp_detail_data = [
        ["Benchmark Source", "Reported Retail Price ($)", "Delta vs Recommended ($)", "Status"]
    ]

    c1 = product.get("comp_1") or product.get("competitor_price", 0.0)
    d1 = rec_price - c1
    comp_detail_data.append(["Competitor 1 (Market Benchmark)", f"${c1:,.2f}", f"{'+' if d1>0 else ''}${d1:,.2f}", "Priced Below" if d1 < 0 else "Priced Above"])

    if product.get("comp_2"):
        c2 = product["comp_2"]
        d2 = rec_price - c2
        comp_detail_data.append(["Competitor 2 (Alternative)", f"${c2:,.2f}", f"{'+' if d2>0 else ''}${d2:,.2f}", "Priced Below" if d2 < 0 else "Priced Above"])

    if product.get("comp_3"):
        c3 = product["comp_3"]
        d3 = rec_price - c3
        comp_detail_data.append(["Competitor 3 (Premium/Discount)", f"${c3:,.2f}", f"{'+' if d3>0 else ''}${d3:,.2f}", "Priced Below" if d3 < 0 else "Priced Above"])

    comp_table = Table(comp_detail_data, colWidths=[180, 120, 120, 120])
    comp_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), primary_color),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, card_bg]),
        ('BOX', (0,0), (-1,-1), 0.5, border_color),
        ('INNERGRID', (0,0), (-1,-1), 0.5, border_color),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(comp_table)
    elements.append(Spacer(1, 10))

    # 5. Embedded Matplotlib Chart
    elements.append(Paragraph("Visual Price Comparison vs. Competitor Benchmarks", section_heading))
    chart_buf = _render_price_comparison_chart_image(product)
    chart_img = Image(chart_buf, width=6.5 * inch, height=2.6 * inch)
    elements.append(chart_img)
    elements.append(Spacer(1, 8))

    # 6. Econometric Methodology Footer
    footer_text = (
        f"<b>PricePilot Econometric Foundation:</b> Price recommendation computed via OLS fixed-effects model "
        f"(beta = {meta['elasticity_coefficient']:.3f}, R² = {meta['r_squared']:.3f}, p = {meta['p_value']:.4f}). "
        f"Optimized to maximize expected catalog gross margin without violating competitive clearing thresholds."
    )
    elements.append(Paragraph(footer_text, ParagraphStyle("Footer", fontName="Helvetica", fontSize=7, leading=9.5, textColor=slate_muted)))

    doc.build(elements)
    buffer.seek(0)
    return buffer

