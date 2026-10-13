from __future__ import annotations

import os
import uuid
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.core.config import settings
from app.services import analytics_service

REPORTS_DIR = os.path.normpath(os.path.join(os.path.dirname(__file__), "..", "..", "..", settings.REPORTS_DIR))


def _collect_report_data(db: Session, filters: dict) -> dict:
    kpis = analytics_service.get_kpis(db, **filters)
    top_products = analytics_service.get_top_products(db, n=10, **filters)
    category_breakdown = analytics_service.get_breakdown(db, by="category", **filters)
    channel_breakdown = analytics_service.get_breakdown(db, by="channel", **filters)
    competitor = analytics_service.get_competitor_analysis(db, product_id=filters.get("product_id"))
    customer = analytics_service.get_customer_analytics(db, product_id=filters.get("product_id"),
                                                          category=filters.get("category"))
    return {
        "kpis": kpis, "top_products": top_products, "category_breakdown": category_breakdown,
        "channel_breakdown": channel_breakdown, "competitor": competitor, "customer": customer,
    }


def generate_report(db: Session, report_type: str, file_format: str, filters: dict) -> str:
    os.makedirs(REPORTS_DIR, exist_ok=True)
    data = _collect_report_data(db, filters)
    report_id = str(uuid.uuid4())[:8]
    filename = f"{report_type}_{report_id}.{file_format}"
    path = os.path.join(REPORTS_DIR, filename)

    if file_format == "pdf":
        _write_pdf(path, report_type, filters, data)
    elif file_format == "xlsx":
        _write_excel(path, report_type, data)
    elif file_format == "csv":
        _write_csv(path, data)
    else:
        raise ValueError(f"Unsupported format: {file_format}")

    return path


def _write_pdf(path: str, report_type: str, filters: dict, data: dict) -> None:
    doc = SimpleDocTemplate(path, pagesize=A4)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"PricePilot AI — {report_type.replace('_', ' ').title()} Report", styles["Title"]),
        Paragraph(f"Generated: {datetime.utcnow().isoformat()} UTC", styles["Normal"]),
        Paragraph(f"Filters: {filters}", styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Executive Summary", styles["Heading2"]),
    ]
    kpis = data["kpis"]
    if kpis.get("has_data"):
        story.append(Paragraph(
            f"Total revenue was ${kpis['total_revenue']:,.2f} across {kpis['total_units_sold']:,} units sold, "
            f"with an estimated gross profit of ${kpis['gross_profit_estimated']:,.2f} "
            f"({kpis['profit_margin_pct_estimated']}% estimated margin). "
            f"Average competitor price gap was {kpis['competitor_price_gap']}.",
            styles["Normal"],
        ))
    story.append(Spacer(1, 12))
    story.append(Paragraph("KPI Summary", styles["Heading2"]))
    kpi_table_data = [["Metric", "Value"]] + [[k, str(v)] for k, v in kpis.items() if k != "has_data"]
    story.append(_styled_table(kpi_table_data))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Top Products by Revenue", styles["Heading2"]))
    top = data["top_products"][:10]
    if top:
        rows = [["Product", "Category", "Revenue", "Units"]] + [
            [p.get("product_name", p.get("product_id")), p.get("category", ""),
             f"${p.get('revenue', 0):,.2f}", str(p.get("units", 0))] for p in top
        ]
        story.append(_styled_table(rows))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Category Breakdown", styles["Heading2"]))
    cat = data["category_breakdown"][:10]
    if cat:
        rows = [["Category", "Revenue", "Units"]] + [
            [c["label"], f"${c['revenue']:,.2f}", str(c["units"])] for c in cat
        ]
        story.append(_styled_table(rows))
    story.append(Spacer(1, 12))

    story.append(Paragraph("AI Note", styles["Heading2"]))
    story.append(Paragraph(
        "All figures above are computed directly from the operational, competitor and customer datasets. "
        "Profitability figures are labelled 'estimated' because they are derived from unit cost x units sold.",
        styles["Normal"],
    ))

    doc.build(story)


def _styled_table(rows: list[list[str]]) -> Table:
    t = Table(rows, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
    ]))
    return t


def _write_excel(path: str, report_type: str, data: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI Summary"
    ws.append(["Metric", "Value"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for k, v in data["kpis"].items():
        if k != "has_data":
            ws.append([k, v])

    ws2 = wb.create_sheet("Top Products")
    if data["top_products"]:
        ws2.append(list(data["top_products"][0].keys()))
        for row in data["top_products"]:
            ws2.append(list(row.values()))

    ws3 = wb.create_sheet("Category Breakdown")
    if data["category_breakdown"]:
        ws3.append(list(data["category_breakdown"][0].keys()))
        for row in data["category_breakdown"]:
            ws3.append(list(row.values()))

    wb.save(path)


def _write_csv(path: str, data: dict) -> None:
    kpi_df = pd.DataFrame([{"metric": k, "value": v} for k, v in data["kpis"].items() if k != "has_data"])
    kpi_df.to_csv(path, index=False)
