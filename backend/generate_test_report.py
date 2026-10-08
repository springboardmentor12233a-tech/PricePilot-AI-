"""
Runs the pytest suite and writes the results to an Excel workbook.

Usage (from the backend/ folder, with your venv active):
    python generate_test_report.py                 # runs tests/ -> PricePilot_AI_Test_Report.xlsx
    python generate_test_report.py tests/ out.xlsx # custom test path / output name

Requires: pytest, openpyxl  (pip install pytest openpyxl)
"""
import sys
import datetime
from collections import OrderedDict

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
STATUS_FILLS = {
    "PASSED": PatternFill("solid", fgColor="C6EFCE"),
    "FAILED": PatternFill("solid", fgColor="FFC7CE"),
    "ERROR": PatternFill("solid", fgColor="FFC7CE"),
    "SKIPPED": PatternFill("solid", fgColor="FFEB9C"),
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class ResultCollector:
    """pytest plugin that records one row per test."""

    def __init__(self):
        self.rows = OrderedDict()   # nodeid -> dict
        self.docs = {}              # nodeid -> docstring

    def pytest_collection_modifyitems(self, items):
        for item in items:
            doc = getattr(getattr(item, "function", None), "__doc__", None)
            self.docs[item.nodeid] = " ".join(doc.split()) if doc else ""

    def pytest_runtest_logreport(self, report):
        row = self.rows.setdefault(
            report.nodeid,
            {"status": None, "duration": 0.0, "message": ""},
        )
        row["duration"] += report.duration
        if report.when == "call":
            if report.passed:
                row["status"] = "PASSED"
            elif report.failed:
                row["status"] = "FAILED"
                row["message"] = _short(report.longreprtext)
            elif report.skipped:
                row["status"] = "SKIPPED"
                row["message"] = _skip_reason(report)
        elif report.when == "setup" and report.failed:
            row["status"] = "ERROR"
            row["message"] = _short(report.longreprtext)
        elif report.when == "setup" and report.skipped:
            row["status"] = "SKIPPED"
            row["message"] = _skip_reason(report)


def _short(text, limit=500):
    text = " ".join(str(text).split())
    return text if len(text) <= limit else text[: limit - 3] + "..."


def _skip_reason(report):
    lr = report.longrepr
    if isinstance(lr, tuple) and len(lr) == 3:
        return _short(lr[2])
    return _short(lr)


def humanise(name):
    name = name.split("[")[0]
    if name.startswith("test_"):
        name = name[5:]
    return name.replace("_", " ").capitalize()


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def build_workbook(collector, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    headers = ["Test ID", "Module", "Test Name", "What it checks", "Status",
               "Duration (s)", "Failure / skip message"]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    for i, (nodeid, r) in enumerate(collector.rows.items(), start=1):
        path, _, test = nodeid.partition("::")
        module = path.split("/")[-1].replace(".py", "")
        desc = collector.docs.get(nodeid) or humanise(test.split("::")[-1])
        ws.append([f"TC-{i:03d}", module, test, desc, r["status"] or "NOT RUN",
                   round(r["duration"], 3), r["message"]])

    last = ws.max_row
    widths = [10, 20, 48, 55, 12, 13, 60]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.font = Font(name=FONT)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        status_cell = row[4]
        status_cell.alignment = Alignment(horizontal="center", vertical="top")
        fill = STATUS_FILLS.get(status_cell.value)
        if fill:
            status_cell.fill = fill
            status_cell.font = Font(name=FONT, bold=True)
        row[5].number_format = "0.000"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"

    # ---- Summary sheet (formulas so it stays in sync with the Test Cases tab)
    sm = wb.create_sheet("Summary", 0)
    sm["A1"] = "PricePilot AI - API & Model Test Report"
    sm["A1"].font = Font(name=FONT, bold=True, size=14)
    sm["A2"] = f"Generated: {datetime.datetime.now():%Y-%m-%d %H:%M}"
    sm["A2"].font = Font(name=FONT, italic=True, color="595959")

    rng = f"'Test Cases'!$E$2:$E${last}"
    rows = [
        ("Total tests", f"=COUNTA({rng})"),
        ("Passed", f'=COUNTIF({rng},"PASSED")'),
        ("Failed", f'=COUNTIF({rng},"FAILED")'),
        ("Errors", f'=COUNTIF({rng},"ERROR")'),
        ("Skipped", f'=COUNTIF({rng},"SKIPPED")'),
        ("Pass rate", "=IF(B4=0,0,B5/B4)"),
    ]
    for offset, (label, formula) in enumerate(rows):
        r = 4 + offset
        sm.cell(row=r, column=1, value=label).font = Font(name=FONT, bold=True)
        c = sm.cell(row=r, column=2, value=formula)
        c.font = Font(name=FONT)
        c.alignment = Alignment(horizontal="right")
    sm["B9"].number_format = "0.0%"

    # per-module breakdown
    start = 12
    mod_headers = ["Module", "Total", "Passed", "Failed / Error", "Pass rate"]
    for c, h in enumerate(mod_headers, start=1):
        sm.cell(row=start, column=c, value=h)
    style_header(sm, start, len(mod_headers))
    modules = list(OrderedDict.fromkeys(
        r.value for r in ws["B"][1:] if r.value))
    mrng = f"'Test Cases'!$B$2:$B${last}"
    for i, mod in enumerate(modules, start=1):
        r = start + i
        sm.cell(row=r, column=1, value=mod)
        sm.cell(row=r, column=2, value=f'=COUNTIF({mrng},A{r})')
        sm.cell(row=r, column=3, value=f'=COUNTIFS({mrng},A{r},{rng},"PASSED")')
        sm.cell(row=r, column=4,
                value=f'=COUNTIFS({mrng},A{r},{rng},"FAILED")+COUNTIFS({mrng},A{r},{rng},"ERROR")')
        sm.cell(row=r, column=5, value=f"=IF(B{r}=0,0,C{r}/B{r})").number_format = "0.0%"
        for c in range(1, 6):
            cell = sm.cell(row=r, column=c)
            cell.font = Font(name=FONT)
            cell.border = BORDER
    sm.column_dimensions["A"].width = 28
    for col in "BCDE":
        sm.column_dimensions[col].width = 16

    wb.save(out_path)


def main():
    test_path = sys.argv[1] if len(sys.argv) > 1 else "tests"
    out_path = sys.argv[2] if len(sys.argv) > 2 else "PricePilot_AI_Test_Report.xlsx"

    collector = ResultCollector()
    exit_code = pytest.main([test_path, "-q", "-p", "no:cacheprovider"],
                            plugins=[collector])
    if not collector.rows:
        print("No tests were collected - check the path:", test_path)
        sys.exit(1)

    build_workbook(collector, out_path)
    passed = sum(1 for r in collector.rows.values() if r["status"] == "PASSED")
    print(f"\nWrote {out_path}: {passed}/{len(collector.rows)} tests passed.")
    sys.exit(int(exit_code))


if __name__ == "__main__":
    main()
